#!/usr/bin/env python3
"""Resolve the declared protection policy into the cells it actually unlocks.

WHY A RESOLVER AND NOT A RANGE LIST. The question "which cells may a user type
in" is already answered, several times over, by the contracts that declare the
inputs: `editable: true` on a Setup input, on a register column, on a Config
master; the year columns of a profiling or inflation grid. A second list of
ranges would be a second answer, and the two would disagree the first time a
column was added. So the policy declares RULES that name the owning contract,
and this is the one place allowed to interpret them.

WHAT IT PRODUCES. Two things from one resolution, so they cannot disagree:

  * the workbook itself   - every resolved cell gets `locked=False`, and Excel's
                            own default leaves everything else locked
  * a projection          - `phase10_protection_inspection.json`, so the Windows
                            acceptance run asks the same question this answered

WHAT IT DOES NOT DO. It applies no protection. Marking a cell unlocked and
protecting a sheet are different acts by different owners: this marks, Stage B
protects, and modProtection re-protects on open.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter

__all__ = ["resolve_unlocked", "apply_protection", "emit_protection_projection"]

UNLOCKED = Protection(locked=False)


def _column_index(letter: str) -> int:
    total = 0
    for char in letter:
        total = total * 26 + (ord(char.upper()) - 64)
    return total


def _table_body(table: Any) -> list[str]:
    """The addresses of a table's USER rows, in every column.

    A SEED ROW IS NOT A USER ROW. The SAR identity on the Setup FX table and the
    seeded currency on Config are model constants that happen to live inside a
    user table. `TableSpec.first_user_row` already draws that line - it is the
    loader's own accessor, not a second opinion - and unlocking a seed row would
    be the single easiest way to let somebody redefine the reporting currency by
    typing over it.
    """
    return [f"{get_column_letter(table.first_col_index + offset)}{row}"
            for row in range(table.first_user_row, table.last_data_row + 1)
            for offset in range(len(table.columns))]


def resolve_unlocked(structure: Any, contract: Any, drivers: Any) -> dict[str, list[str]]:
    """`{sheet: [address, ...]}` for every sheet the policy names.

    A sheet whose `unlock` list is empty resolves to an EMPTY LIST rather than
    being absent: "nothing is unlocked here" is a decision, and the projection
    has to be able to show that it was made.
    """
    policy = structure.protection
    if not policy:
        return {}
    resolved: dict[str, list[str]] = {}
    for entry in policy["sheets"]:
        sheet = str(entry["sheet"])
        addresses: list[str] = []
        for rule in entry.get("unlock") or []:
            addresses.extend(_resolve_rule(rule, sheet, structure, contract, drivers))
        # Deterministic and duplicate-free: two rules may legitimately reach the
        # same cell, and the projection must not depend on which ran first.
        resolved[sheet] = sorted(set(addresses), key=_address_sort_key)
    return resolved


def _address_sort_key(address: str) -> tuple[int, int]:
    letters = "".join(c for c in address if c.isalpha())
    digits = "".join(c for c in address if c.isdigit())
    return (int(digits), _column_index(letters))


def _resolve_rule(rule: dict[str, Any], sheet: str, structure: Any,
                  contract: Any, drivers: Any) -> list[str]:
    name = str(rule["rule"])

    if name == "editable_inputs":
        return [spec.cell for spec in contract.inputs.values()
                if spec.sheet == sheet and spec.editable]

    if name == "editable_input_table":
        table = contract.tables[str(rule["key"])]
        return _table_body(table) if table.editable else []

    if name == "editable_config_tables":
        return [address for table in contract.config_tables if table.editable
                for address in _table_body(table)]

    if name == "editable_register_columns":
        register = drivers.register_for_sheet(sheet)
        return [f"{get_column_letter(register.first_col_index + offset)}{row}"
                for offset, column in enumerate(register.columns)
                if column.editable
                for row in range(register.first_data_row, register.last_data_row + 1)]

    if name == "grid_year_columns":
        grid = structure.grids[str(rule["key"])]
        # THE FIXED COLUMNS STAY LOCKED. A permanent id, a synchronised profile
        # name and a trace copy of a description are all model-controlled; the
        # grid's own note says editing the trace copy changes nothing, and a
        # cell that changes nothing should not invite the attempt.
        # `first_year_column_index` is a ZERO-BASED OFFSET from the grid's first
        # column, not an absolute index - the loader's own docstring says so.
        # Reading it as absolute unlocked the fixed columns, which is precisely
        # the permanent-id column this policy exists to keep locked.
        start = grid.first_col_index + grid.first_year_column_index()
        return [f"{get_column_letter(start + offset)}{row}"
                for offset in range(_year_column_count(structure))
                for row in range(grid.first_data_row, grid.last_data_row + 1)]

    raise ValueError(f"unknown protection rule {name!r}")


def _year_column_count(structure: Any) -> int:
    """The widest project-year span the grids can carry.

    Unlocking the DECLARED MAXIMUM rather than the currently applied duration is
    deliberate: the applied duration changes when the timeline is applied, and a
    protected workbook whose newly generated year columns were locked would
    refuse the very input it had just created room for.
    """
    return int(structure.limits.max_generated_year_columns)


def apply_protection(workbook: Any, unlocked: dict[str, list[str]]) -> None:
    """Mark exactly the resolved cells unlocked, and nothing else.

    Every other cell keeps openpyxl's default `locked=True`, so this never has to
    lock anything: the absence of a rule IS the lock.
    """
    for sheet, addresses in unlocked.items():
        worksheet = workbook[sheet]
        for address in addresses:
            worksheet[address].protection = UNLOCKED


def emit_protection_projection(path: Path, structure: Any,
                               unlocked: dict[str, list[str]]) -> dict[str, Any]:
    policy = structure.protection
    projection = {
        "passwordless": bool(policy["passwordless"]),
        "user_interface_only": bool(policy["user_interface_only"]),
        "protect_structure": bool(policy["protect_structure"]),
        "protect_windows": bool(policy["protect_windows"]),
        "sheets": [
            {
                "sheet": sheet,
                "protect": True,
                "unlocked_count": len(addresses),
                "unlocked": addresses,
            }
            for sheet, addresses in unlocked.items()
        ],
    }
    path.write_text(json.dumps(projection, indent=2) + "\n", encoding="utf-8")
    return projection
