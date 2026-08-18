"""Render the Phase-5 calculation workspace onto `_Calc`.

REPRESENTATION, NOT CALCULATION. This module projects the physical shape declared
by `spec/calc_contract.yaml` — two scalar blocks and five ListObjects — and writes
nothing that a calculation would later produce. It never imports `calc_oracle`,
never computes a factor, a total or a fingerprint, and never writes a formula.

--------------------------------------------------------------------------------
WHY THE WORKSPACE SHIPS EMPTY
--------------------------------------------------------------------------------
`_Calc` is a WRITTEN RECORD of an in-memory kernel, not a spreadsheet that
computes. A freshly built workbook has had no calculation, so:

  * every `calc_state` value cell is blank except the two that record the ABSENCE
    of a calculation — attempt result `NONE` and derived status `NOT CALCULATED`;
  * every `calc_totals` value cell is blank. **Blank is not zero.** Blank means no
    analytical calculation has committed; zero means a calculated total of zero.
    Seeding zeros would make those two states indistinguishable forever;
  * `Fingerprint Version` (C15) is blank even though `FP_VERSION` is 1 in the
    contract. The version stamp records which algorithm produced the STORED
    digest, and there is no stored digest until a successful commit writes both;
  * the five tables carry their headers and no semantic row. There is no project
    year, no referenced currency, no driver and no annual row, because those are
    calculation outputs.

--------------------------------------------------------------------------------
PHASE-4 TERRITORY IS NOT TOUCHED
--------------------------------------------------------------------------------
Rows 1-11 of `_Calc` belong to Phase 4 — the permanent-ID counters at C10 and C11
and their labels and notes. The contract declares that reservation and the loader
proves no Phase-5 block intersects it. This module starts at row 13 and never
writes, clears, restyles or resizes through anything above it.
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .calc_loader import CalcContract, CalcTable, ScalarBlock
from .styling import StyleBook

PLACEHOLDER_ROW_OFFSET = 1
"""The data body of a Phase-5 ListObject is exactly one physically blank row.

WHY NOT A HEADER-ONLY TABLE. A `ref` that spans only the header row describes a
ListObject with zero data rows. Excel has no such object: deleting the last row of
a table leaves one blank row rather than none, and a single-row `ref` with
`headerRowCount=1` is the shape Excel treats as damaged and offers to repair. A
workbook that prompts for repair on first open is not an acceptable Stage-A
artifact, so the minimum PHYSICALLY valid body is used instead.

The placeholder is semantically nothing: no value, no formula, no identifier, no
zero. Post-build verification asserts it is blank and counts the workspace as
having zero calculation rows, and the runtime that later adds a real row will
overwrite it exactly as the Phase-3 and Phase-4 registers already do with their
own reserved blank rows.
"""


def render_calc_workspace(
    worksheet: Worksheet, calc: CalcContract, styles: StyleBook
) -> None:
    """The whole Phase-5 workspace: both scalar blocks and all five tables."""
    _render_scalar_block(worksheet, calc.calc_state, styles)
    _render_scalar_block(worksheet, calc.calc_totals, styles)
    for table in calc.all_tables:
        _render_table(worksheet, table, styles)


def _render_scalar_block(
    worksheet: Worksheet, block: ScalarBlock, styles: StyleBook
) -> None:
    """Labels, value cells and notes for one declared scalar block.

    The value cell carries its contract number format whether or not it carries a
    value, so a later runtime write lands in an already correctly formatted cell
    and cannot change the audit presentation by writing.
    """
    for field in block.fields:
        label = worksheet[f"{block.label_column}{field.row}"]
        label.value = field.label
        label.font = styles.label

        cell = worksheet[f"{block.value_column}{field.row}"]
        # `initial` is None for every cell that must ship BLANK. Assigning None to
        # an openpyxl cell leaves it genuinely empty - not "", not 0.
        cell.value = field.initial
        cell.number_format = field.number_format
        styles.apply_locked(cell)

        if field.note:
            note = worksheet[f"{block.note_column}{field.row}"]
            note.value = field.note
            note.font = styles.note


def _render_table(worksheet: Worksheet, table: CalcTable, styles: StyleBook) -> None:
    """One Phase-5 ListObject: header row, formats, and one blank placeholder row."""
    body_row = table.header_row + PLACEHOLDER_ROW_OFFSET

    for index, column in enumerate(table.columns):
        letter = get_column_letter(table.first_column_index + index)

        header = worksheet[f"{letter}{table.header_row}"]
        header.value = column.header
        styles.apply_table_header(header)

        # The placeholder body cell is formatted and locked but left empty. Its
        # number format is the column's, so nothing about presentation depends on a
        # calculation having run.
        cell = worksheet[f"{letter}{body_row}"]
        cell.number_format = column.number_format
        styles.apply_locked(cell)

    excel_table = Table(displayName=table.table_name, ref=table_ref(table))
    excel_table.tableStyleInfo = TableStyleInfo(
        name=styles.table_style_name,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    worksheet.add_table(excel_table)


def table_ref(table: CalcTable) -> str:
    """The OOXML `ref` of one Phase-5 ListObject: header row plus one blank row.

    Exposed so the verifier can assert the PHYSICAL reference in the generated
    artifact against the same derivation, rather than against a literal restated in
    a second place.
    """
    body_row = table.header_row + PLACEHOLDER_ROW_OFFSET
    return (
        f"{table.first_column}{table.header_row}:{table.last_column}{body_row}"
    )


def placeholder_row(table: CalcTable) -> int:
    """The single physically blank data row of *table*."""
    return table.header_row + PLACEHOLDER_ROW_OFFSET
