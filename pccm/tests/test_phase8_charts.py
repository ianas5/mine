#!/usr/bin/env python3
"""Phase 8, Step 3: the analytical chart layer, proved against its sources.

WHAT THIS SLICE IS. Four charts on the Dashboard - a cumulative cost curve, a
total-cost histogram, an annual cash flow and a driver tornado - each pointed at
a bridge block on Results, and each bridge block derived from a surface that was
already accepted. No engine, no statistic, no ranking and no VBA.

WHAT THESE CONTROLS EXIST TO CATCH. A chart fails silently in ways a table does
not. Excel plots a zero-length string as ZERO, so a series over a 200-row window
draws 196 fabricated years of a project costing nothing, and it looks like data.
A tornado re-sorted by signed rho looks tidier and is wrong. A histogram whose
last bin is half-open loses the maximum and nobody counts. So the controls below
check the SHAPE OF THE ABSENCE as hard as they check the presence: NA() where
there is no point, count conservation across the bins, the rank order taken and
not re-taken.

WHAT THEY CANNOT PROVE: there is no Excel here. Nothing below claims a formula
was evaluated, a chart was drawn, or a series rendered any particular way.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

PCCM_ROOT = Path(__file__).resolve().parent.parent
REPO_ROOT = PCCM_ROOT.parent
sys.path.insert(0, str(PCCM_ROOT / "builder"))

import pytest  # noqa: E402
import yaml  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from pccm_builder.phase8_charts import (  # noqa: E402
    build_phase8_charts_inspection,
    validate_phase8_charts_inspection,
)
from pccm_builder.spec_loader import SpecError, load_spec  # noqa: E402
from pccm_builder.structure_loader import load_structure_contract  # noqa: E402

SPEC = PCCM_ROOT / "spec"
BUILD = PCCM_ROOT / "build"
SRC = PCCM_ROOT / "src" / "vba"
MANIFEST = SPEC / "workbook.yaml"

# THE COMMITS THE EARLIER STEPS WERE ACCEPTED AT. Their geometry is surface a
# Windows run has been produced against; P8-3 may extend the manifest but may
# not move a byte of either block.
P81_ACCEPTANCE = "35bd6ce"
P82_ACCEPTANCE = "7ff5dc7"

_CACHE: dict = {}


def _shell() -> dict:
    if "shell" not in _CACHE:
        _CACHE["shell"] = yaml.safe_load(MANIFEST.read_text(encoding="utf-8"))["phase6_shell"]
    return _CACHE["shell"]


def _sim_contract():
    from pccm_builder import load_sim_contract
    return load_sim_contract(SPEC / "sim_contract.yaml")


def _projection() -> dict:
    if "charts" not in _CACHE:
        _CACHE["charts"] = json.loads(
            (BUILD / "phase8_charts_inspection.json").read_text(encoding="utf-8"))
    return _CACHE["charts"]


def _results_projection() -> dict:
    if "results" not in _CACHE:
        _CACHE["results"] = json.loads(
            (BUILD / "phase8_results_inspection.json").read_text(encoding="utf-8"))
    return _CACHE["results"]


def _workbook():
    if "workbook" not in _CACHE:
        manifest = json.loads((BUILD / "stage_b_manifest.json").read_text(encoding="utf-8"))
        _CACHE["workbook"] = load_workbook(BUILD / manifest["stage_a_filename"])
    return _CACHE["workbook"]


def _results():
    return _workbook()["Results"]


def _dashboard():
    return _workbook()[_projection()["chart_sheet"]]


def _charts() -> list:
    return list(_dashboard()._charts)


def _by_key() -> dict[str, dict]:
    return {chart["key"]: chart for chart in _projection()["charts"]}


def _title(chart) -> str:
    return chart.title.tx.rich.p[0].r[0].t


def _formula(address: str) -> str:
    value = _results()[address].value
    assert isinstance(value, str) and value.startswith("="), (
        f"Results!{address} holds {value!r}, not a formula")
    return value


def _block_formulas(block: str, key: str) -> list[str]:
    """Every formula in one bridge column, top to bottom."""
    entry = _projection()["bridge"][block]
    column = next(c["column"] for c in entry["columns"] if c["key"] == key)
    return [_formula(f"{column}{row}")
            for row in range(int(entry["first_row"]), int(entry["last_row"]) + 1)]


def _git(*args: str) -> str:
    import subprocess
    return subprocess.run(["git", *args], cwd=REPO_ROOT, check=True,
                          stdout=subprocess.PIPE, text=True).stdout


# ===========================================================================
# PRODUCTION MODIFICATIONS DECLARED SINCE THE ACCEPTANCE COMMIT
# ===========================================================================
# THE ORIGINAL CLAIM WAS "no pccm/src change since the acceptance commit", and
# it was true when written. The P8-3 pre-Windows correction made it false on
# purpose: the chart layer needed a live simulation state, and the only honest
# place for it was a fifth adapter in the module that already owns the
# worksheet-safe presentation surface.
#
# SO THE RULE BECOMES DECLARATION, NOT PROHIBITION - the same settlement the
# Phase-7 closure record reached for the same reason. An undeclared change still
# fails. A deletion still fails, declared or not. And a DECLARED file may only
# be ADDED to: the diff against the acceptance commit must remove no line, so a
# declaration cannot be used to quietly edit an accepted procedure.
#
# That last clause is why this is stricter than what it replaces, not weaker.
# The old control could only say "nothing changed"; this one says "exactly this
# changed, additively, and here is why".
# AND ONE MORE THING A DECLARATION HAS TO SAY, added when P8-3's fourth Windows
# run found a REAL DEFECT rather than a gap: WHICH LINES IT MAY REMOVE.
#
# The rule above is additive-only, and that is right for an adapter being added
# to a module. It cannot express a one-line OWNERSHIP correction: reading the
# wrong column is fixed by not reading the wrong column, and there is no way to
# do that additively. Refusing the correction would leave the defect; relaxing
# the rule to "declared files may be edited" would retire the protection.
#
# SO A DECLARATION NAMES ITS REMOVALS EXACTLY. Each entry is (reason, removals):
# an empty tuple keeps the file additive-only, and any other removed line still
# fails. That is stricter than what it replaces, not weaker - the old control
# could say "this file may grow"; this one says "this file may grow, and may
# lose exactly this line".
DECLARED_PRODUCTION_CORRECTIONS = {
    "pccm/src/vba/modResultsState.bas": (
        "P8-3 pre-Windows correction: adds the thin volatile adapter "
        "PCCM_ResultsSimulationState, delegating to the accepted pure evaluator "
        "modSimReport.SimReportDerivedStatus. The chart layer needed a live "
        "SIMULATION state: the four annual adapters answer about the annual "
        "product and read NOT PRODUCED whenever the annual step has not run, and "
        "the persisted (last evaluated) row was proved live at P8-1 to keep "
        "reading CURRENT after a request change. P9-2 adds a SIXTH adapter to "
        "the same module on the same pattern, PCCM_ModelCheckCalculationState, "
        "delegating to modCalcReport.CalcReportDerivedStatus: Model Check needed "
        "a live CALCULATION state, the persisted _Calc C19 row is last-evaluated "
        "and goes on reading CURRENT after an input change, and the owner's own "
        "entry point persists and so may not be called from a cell. Both "
        "additions are whole new procedures; no existing one is touched.",
        (),
    ),
    "pccm/src/vba/modCalcReport.bas": (
        "P9-2: the live calculation state is read-only. Adds "
        "CalcReportDerivedStatus, a public delegation returning the existing "
        "private DeriveStatus over the existing private PrepareCurrentCalculation. "
        "PCCM_CalculationStatus derives the same answer AND writes C19:C20 through "
        "WriteStatusBlock, and Excel forbids a function a worksheet cell called to "
        "change the workbook - the P8-1 defect exactly, one module along. The "
        "derivation was already pure and the persistence is the caller's, so this "
        "exposes the pure half. No derivation, no persistence, no state word and "
        "no existing procedure changed: PCCM_CalculationStatus still derives and "
        "still writes, exactly as Phase 5 accepted it.",
        (),
    ),
    "pccm/src/vba/modSimPostReport.bas": (
        "P8-3 Windows run 4: a risk publishes its RISK NAME. DriverNameOf read "
        "COL_RISK_REGISTER_DESCRIPTION for a risk - the column driver_contract "
        "declares required: false, an optional note - so a risk without one "
        "published no name at all, and the empty string reached the sheet "
        "through a .Value2 array write as a numeric zero. It now reads "
        "COL_RISK_REGISTER_RISK_NAME, which that contract declares required: "
        "true. The cost-line branch is untouched and was always right: a cost "
        "line has no name column and its description is its required label. No "
        "driver id, rank, signed rho, absolute rho, direction, status, ordering, "
        "fingerprint or replay mathematics is touched.",
        ("    column = COL_RISK_REGISTER_DESCRIPTION",),
    ),
}


def _declared_production_changes(git, since: str) -> None:
    """Every production change since *since* is declared, additive, and real.

    THREE SEPARATE FAILURES, because they are three different mistakes:

      undeclared   a module changed and nobody said so - the thing this control
                   has always existed to catch
      deleted      a module the accepted evidence ran against stopped existing,
                   which no declaration may permit
      edited       a DECLARED file lost a line. A declaration buys the right to
                   ADD to a module, never to rewrite what was accepted in it.

    And a fourth, in the other direction: a declaration naming a file that was
    never touched would let the next real change hide beside it.
    """
    changes = git("diff", "--name-status", f"{since}..HEAD", "--", "pccm/src")
    modified, deleted = [], []
    for line in changes.splitlines():
        if not line.strip():
            continue
        state, path = line.split("\t", 1)[0].strip(), line.split("\t", 1)[1].strip()
        if state.startswith("A"):
            continue
        (deleted if state.startswith("D") else modified).append(path)

    assert not deleted, (
        f"a module the accepted evidence ran against was removed after {since}: "
        f"{deleted}")
    undeclared = [p for p in modified if p not in DECLARED_PRODUCTION_CORRECTIONS]
    assert not undeclared, (
        f"production changed after {since} without being declared: {undeclared}")

    for path in modified:
        # ADDITIVE, OR REMOVING EXACTLY WHAT THE DECLARATION NAMED. A declaration
        # is permission to extend a module and - only where it says so, line for
        # line - to correct one. Anything else removed means an accepted
        # procedure was rewritten.
        reason, allowed = DECLARED_PRODUCTION_CORRECTIONS[path]
        removed = [line[1:] for line in
                   git("diff", f"{since}..HEAD", "--", path).splitlines()
                   if line.startswith("-") and not line.startswith("---")]
        undeclared_removals = [line for line in removed if line not in allowed]
        assert not undeclared_removals, (
            f"{path} removes {len(undeclared_removals)} line(s) its declaration "
            f"does not name: {undeclared_removals[:3]}")
        # AND A NAMED REMOVAL THAT NEVER HAPPENED is a licence held in reserve.
        unused = [line for line in allowed if line not in removed]
        assert not unused, (
            f"{path} declares it removes {unused}, and it does not")
        assert len(reason) > 80, f"the declaration for {path} explains nothing"
    # A DECLARATION FOR AN UNTOUCHED FILE IS DECORATION, and the next real
    # change would hide beside it.
    stale = [p for p in DECLARED_PRODUCTION_CORRECTIONS if p not in modified]
    assert not stale, f"declared corrections that changed nothing: {stale}"


# ===========================================================================
# A. GENERAL - THE INVENTORY, AND WHAT IT MAY NOT DISTURB
# ===========================================================================
def test_01_the_workbook_still_has_exactly_fourteen_sheets() -> None:
    """Charts are objects anchored over cells. They need no sheet of their own,
    and a hidden chart-data sheet would be a fifteenth."""
    assert len(_workbook().sheetnames) == 14, _workbook().sheetnames


def test_02_exactly_the_four_projected_charts_exist() -> None:
    expected = _by_key()
    assert sorted(expected) == ["annual_cash_flow", "histogram", "s_curve", "tornado"]
    assert len(_charts()) == 4, len(_charts())
    titles = {_title(chart) for chart in _charts()}
    assert titles == {chart["title"] for chart in expected.values()}
    # AND NOWHERE ELSE IN THE WORKBOOK.
    for worksheet in _workbook().worksheets:
        if worksheet.title == _projection()["chart_sheet"]:
            continue
        assert not list(getattr(worksheet, "_charts", [])), worksheet.title
        assert not list(getattr(worksheet, "_images", [])), worksheet.title


def test_03_no_chart_is_three_dimensional_or_carries_junk() -> None:
    """A THIRD DIMENSION CARRIES NO DATA HERE and distorts the comparison the
    chart exists to make. A legend for one series names what the title already
    said. Category gridlines measure nothing."""
    for chart in _charts():
        assert type(chart).__name__ in ("LineChart", "BarChart"), type(chart).__name__
        assert "3D" not in type(chart).__name__
        assert chart.x_axis.majorGridlines is None
        assert chart.y_axis.majorGridlines is None
        assert chart.x_axis.delete is False and chart.y_axis.delete is False
        expected_legend = len(chart.series) > 1
        assert (chart.legend is not None) == expected_legend, (
            f"{_title(chart)} has a legend that does not match its series count")


def test_04_every_chart_sits_inside_the_reserved_region_and_below_the_summary() -> None:
    """THE P8-2 EXECUTIVE SUMMARY IS ABOVE AND STAYS ABOVE. A chart anchored one
    row too high would sit on top of a number somebody reads."""
    dashboard = _shell()["dashboard"]
    region = dashboard["chart_region"]
    summary_last = max(int(section["first_row"]) + len(section["rows"]) - 1
                       for section in dashboard["sections"])
    assert summary_last < int(region["first_row"])
    seen = set()
    for chart in _charts():
        row = chart.anchor._from.row + 1
        column = chart.anchor._from.col + 1
        assert int(region["first_row"]) <= row <= int(region["last_row"]), row
        assert row > summary_last
        assert (row, column) not in seen, "two charts share an anchor"
        seen.add((row, column))
    assert len(seen) == 4


def test_05_the_reserved_region_cells_are_still_empty() -> None:
    """A CHART IS AN OBJECT OVER CELLS, not a value in them. Anything written
    into the region would be a caption competing with the chart's own title."""
    region = _shell()["dashboard"]["chart_region"]
    sheet = _dashboard()
    for row in range(int(region["first_row"]), int(region["last_row"]) + 1):
        for column in "ABCDEFGH":
            value = sheet[f"{column}{row}"].value
            assert value is None, f"Dashboard!{column}{row} holds {value!r}"


def test_06_no_dashboard_chart_reads_machine_storage() -> None:
    """THE BRIDGE IS THE ONLY THING A CHART MAY SEE, because the bridge is where
    "" became NA(). A chart pointed at `_SimData` would plot the machine's
    blanks as zeros and would go stale the way P7-4's address did."""
    bridge_sheet = _projection()["bridge_sheet"]
    for chart in _charts():
        references = [series.val.numRef.f for series in chart.series]
        references += [series.cat.numRef.f for series in chart.series
                       if series.cat is not None and series.cat.numRef is not None]
        references += [series.cat.strRef.f for series in chart.series
                       if series.cat is not None and series.cat.strRef is not None]
        for reference in references:
            assert reference.startswith(f"'{bridge_sheet}'!"), reference
            for machine in ("_SimData", "_Calc"):
                assert machine not in reference, reference


def test_07_chart_geometry_matches_the_projection_exactly() -> None:
    """ONE AUTHORITY. Type, anchor, size, title, every series range and the
    category range - all of it answerable from the projection, and all of it
    what the workbook actually got."""
    expected = _by_key()
    by_title = {_title(chart): chart for chart in _charts()}
    kinds = {"line": ("LineChart", None), "column": ("BarChart", "col"),
             "bar": ("BarChart", "bar")}
    for spec in expected.values():
        chart = by_title[spec["title"]]
        klass, bar_type = kinds[spec["kind"]]
        assert type(chart).__name__ == klass, spec["key"]
        if bar_type is not None:
            assert chart.type == bar_type, spec["key"]
        column = re.match(r"([A-Z]+)(\d+)", spec["anchor"]).group(1)
        row = int(re.match(r"([A-Z]+)(\d+)", spec["anchor"]).group(2))
        assert chart.anchor._from.row + 1 == row
        assert chart.anchor._from.col + 1 == sum(
            (ord(ch) - 64) * 26 ** index
            for index, ch in enumerate(reversed(column)))
        extent = chart.anchor.ext
        assert round(extent.cx / 360000, 1) == round(spec["width_cm"], 1)
        assert round(extent.cy / 360000, 1) == round(spec["height_cm"], 1)
        actual_series = [series.val.numRef.f for series in chart.series]
        assert actual_series == [
            f"'{_projection()['bridge_sheet']}'!"
            + series["range"].split("!", 1)[1] for series in spec["series"]], spec["key"]
        for series in chart.series:
            reference = (series.cat.numRef.f if series.cat.numRef is not None
                         else series.cat.strRef.f)
            assert reference == (f"'{_projection()['bridge_sheet']}'!"
                                 + spec["categories"]["range"].split("!", 1)[1])


def test_08_every_bridge_block_declares_its_kind_and_its_authority() -> None:
    """DERIVED CHART-ONLY DATA OR A MIRROR - and which one it is has to be
    written down. An unlabelled helper block is an accidental second analytical
    authority waiting to be cited."""
    for name, block in _projection()["bridge"].items():
        if not isinstance(block, dict):
            continue
        assert block["kind"] in ("derived_chart_only", "mirrored"), (name, block["kind"])
        assert len(block["authority"]) > 20, name
    assert _projection()["bridge"]["status"]["kind"] == "mirrored"
    for name in ("annual", "distribution", "drivers"):
        assert _projection()["bridge"][name]["kind"] == "derived_chart_only"


def test_09_the_bridge_does_not_touch_the_accepted_p8_1_cells() -> None:
    """THOSE CELLS CARRY EVIDENCE A WINDOWS RUN WAS PRODUCED AGAINST. An overlap
    would let both sets read whatever was written last."""
    from pccm_builder.verify import _phase6_formula_cells

    spec = load_spec(MANIFEST)
    structure = load_structure_contract(SPEC / "structure_contract.yaml")
    permitted = _phase6_formula_cells(spec, structure)["Results"]
    bridge_rows = set()
    for name, block in _projection()["bridge"].items():
        if not isinstance(block, dict):
            continue
        bridge_rows |= set(range(int(block["first_row"]),
                                 int(block.get("last_row", block["first_row"])) + 1))
    accepted = _results_projection()
    accepted_rows = (
        {int(f["row"]) for f in accepted["run_stamp"]["fields"]}
        | {int(m["row"]) for m in accepted["summary"]["metrics"]}
        | set(accepted["selected"].values())
        | {int(v["row"]) for v in accepted["state"].values()}
        | set(range(int(accepted["annual"]["first_row"]),
                    int(accepted["annual"]["first_row"])
                    + int(accepted["annual"]["row_window"])))
        | set(accepted["reconciliation"]["rows"].values()))
    assert not (bridge_rows & accepted_rows), sorted(bridge_rows & accepted_rows)
    assert permitted, "the Results permitted set is empty"


def test_10_no_vba_and_no_new_state_algorithm_arrived_with_the_charts() -> None:
    """THE CHART LAYER ITSELF IS FORMULA-ONLY, and the state words still belong
    to their owner.

    RESTATED AT THE P8-3 PRE-WINDOWS CORRECTION. This read "no production VBA
    changed", which was true of the chart layer and became false when the
    correction added the live simulation adapter the charts needed. The claim is
    now DECLARATION rather than prohibition - and it is stricter, because a
    declared file must also be purely additive."""
    _declared_production_changes(_git, P82_ACCEPTANCE)
    for module in sorted(SRC.glob("*.bas")):
        code = "\n".join(line for line in module.read_text(encoding="utf-8").splitlines()
                         if not line.lstrip().startswith("'"))
        for word in ("Chart", "Histogram", "Tornado", "SCurve"):
            assert word not in code, f"{module.name} names {word}"


def test_11_the_accepted_p8_1_and_p8_2_geometry_did_not_move() -> None:
    """P8-3 MAY EXTEND THE MANIFEST; it may not move a byte of a surface a
    Windows run was produced against. Read from the two acceptance commits
    rather than asserted."""
    p81 = yaml.safe_load(_git("show", f"{P81_ACCEPTANCE}:pccm/spec/workbook.yaml"))
    assert p81["phase6_shell"]["results"] == _shell()["results"], (
        "the accepted P8-1 Results block changed in this step")
    p82 = yaml.safe_load(_git("show", f"{P82_ACCEPTANCE}:pccm/spec/workbook.yaml"))
    accepted = p82["phase6_shell"]["dashboard"]
    current = _shell()["dashboard"]
    for key in ("sheet", "source_sheet", "label_column", "nominal_column",
                "pv_column", "mirror_formula", "number_formats"):
        assert accepted[key] == current[key], f"the P8-2 dashboard {key} moved"
    # THE ACCEPTED SECTIONS ARE UNCHANGED; P8-3 ADDS ONE AND MOVES THE REGION.
    assert current["sections"][:len(accepted["sections"])] == accepted["sections"], (
        "an accepted P8-2 dashboard section changed")
    added = current["sections"][len(accepted["sections"]):]
    assert [section["key"] for section in added] == ["chart_status"], added


# ===========================================================================
# B. THE S-CURVE
# ===========================================================================
def _annual_guard(year: int) -> str:
    annual = _shell()["results"]["annual"]
    nominal = _shell()["results"]["nominal_column"]
    return (f'OR(${nominal}${annual["distribution_state_row"]}="NOT PRODUCED",'
            f'{year}>${nominal}${annual["year_count_row"]})')


def test_20_the_cumulative_series_is_a_running_sum_of_the_published_profile() -> None:
    """NOTHING IS RE-ALLOCATED. The S-curve is a running SUM of the annual
    profile rows the accepted P8-1 table already publishes - display arithmetic
    over cells on the same sheet, not a second stochastic allocation."""
    results = _shell()["results"]["annual"]
    first = int(results["first_row"])
    display = {c["field"]: c["column"] for c in results["columns"]}
    for measure, column in (("nominal", display["nominal"]), ("pv", display["pv"])):
        formulas = _block_formulas("annual", f"cumulative_{measure}")
        for offset, formula in enumerate(formulas):
            span = f"SUM(${column}${first}:${column}${first + offset})"
            assert span in formula, (offset, measure, formula)
            assert _annual_guard(offset + 1) in formula, (offset, formula)
    # AND THE TWO MEASURES ARE NOT SWAPPED.
    nominal = _block_formulas("annual", "cumulative_nominal")[0]
    pv = _block_formulas("annual", "cumulative_pv")[0]
    assert f"${display['nominal']}$" in nominal and f"${display['pv']}$" not in nominal
    assert f"${display['pv']}$" in pv and f"${display['nominal']}$" not in pv


def test_21_the_annual_series_is_bounded_by_the_stamped_year_count() -> None:
    """NO POST-WINDOW ZERO YEARS. Every row past the stamped count returns NA(),
    which every chart type declines to plot. `""` would be drawn as zero - a
    fabricated year of a project costing nothing - which is the exact reason
    this bridge exists rather than the chart reading the table."""
    for key in ("project_index", "calendar_year", "annual_nominal",
                "cumulative_nominal", "cumulative_pv"):
        formulas = _block_formulas("annual", key)
        assert len(formulas) == _projection()["bridge"]["annual"]["row_count"]
        for offset, formula in enumerate(formulas):
            assert formula.startswith("=IF("), formula
            assert ",NA()," in formula, (key, offset, formula)
            assert _annual_guard(offset + 1) in formula, (key, offset)
            assert '""' not in formula, (
                f"{key} row {offset} can return a blank a chart would plot as zero")


def test_22_the_s_curve_plots_the_cumulative_series_against_calendar_year() -> None:
    spec = _by_key()["s_curve"]
    assert spec["kind"] == "line"
    assert [series["key"] for series in spec["series"]] == [
        "cumulative_nominal", "cumulative_pv"]
    assert spec["categories"]["key"] == "calendar_year"
    assert spec["source_block"] == "annual"
    assert spec["state_source"] == "profile_state"
    # A LINE CHART OF A CUMULATIVE SERIES IS NOT SMOOTHED. Interpolating between
    # two published years would draw money that was never allocated.
    for chart in _charts():
        if _title(chart) != spec["title"]:
            continue
        for series in chart.series:
            assert series.smooth is False or series.smooth is None


# ===========================================================================
# C. THE HISTOGRAM
# ===========================================================================
def test_30_the_bin_contract_is_deterministic_and_projected() -> None:
    """A FIXED COUNT, NOT A RULE THAT READS THE DATA. The same run always bins
    the same way, so two readers never see two histograms of one result."""
    contract = _projection()["bin_contract"]
    assert contract["bin_count"] == _projection()["bridge"]["distribution"]["row_count"]
    assert contract["bin_count"] >= 2
    assert contract["measure"] == "nominal"
    for clause in ("equal width", "[lower, upper)", "<= upper", "NA()"):
        assert clause in json.dumps(contract), clause


def test_31_the_edges_come_from_the_published_minimum_and_maximum() -> None:
    """NO STATISTIC IS RECOMPUTED. The extremes are the ones the run published
    on Results; a bridge that scanned the iteration column for its own min and
    max would be a second implementation of two numbers already settled."""
    summary = {m["key"]: m["row"] for m in _results_projection()["summary"]["metrics"]}
    nominal = _results_projection()["columns"]["nominal"]
    minimum = f"${nominal}${summary['minimum']}"
    maximum = f"${nominal}${summary['maximum']}"
    lowers = _block_formulas("distribution", "lower")
    uppers = _block_formulas("distribution", "upper")
    count = _projection()["bin_contract"]["bin_count"]
    for index, (lower, upper) in enumerate(zip(lowers, uppers)):
        width = f"(({maximum}-{minimum})/{count})"
        assert f"({minimum}+{index}*{width})" in lower, (index, lower)
        assert f"({minimum}+{index + 1}*{width})" in upper, (index, upper)
    assert "MIN(" not in "".join(lowers) and "MAX(" not in "".join(uppers)
    assert "PERCENTILE" not in "".join(lowers + uppers)


def test_32_every_iteration_lands_in_exactly_one_bin() -> None:
    """COUNT CONSERVATION, READ OFF THE COMPARATORS. Bins are half-open so no
    iteration is counted twice, and the LAST one closes so the maximum is
    counted at all rather than falling off the end."""
    counts = _block_formulas("distribution", "count")
    for index, formula in enumerate(counts[:-1]):
        assert '"<"&' in formula, (index, "not half-open")
        assert '"<="&' not in formula, (index, "double-counts its upper edge")
        assert '">="&' in formula
    assert '"<="&' in counts[-1], "the last bin drops the maximum"
    assert '">="&' in counts[-1]


def test_33_an_unpublished_or_degenerate_run_fabricates_no_bin() -> None:
    """A HISTOGRAM OF ZEROS IS A CONFIDENT PICTURE OF A RUN THAT NEVER HAPPENED.
    And a run whose iterations all produced the same total has no width to bin:
    the whole population goes in the first bin and the rest return NA()."""
    stamp = {f["key"]: f["row"] for f in _results_projection()["run_stamp"]["fields"]}
    nominal = _results_projection()["columns"]["nominal"]
    published = f"${nominal}${stamp['run_id']}"
    counts = _block_formulas("distribution", "count")
    for index, formula in enumerate(counts):
        assert formula.startswith(f'=IF({published}="",NA(),'), (index, formula)
        if index == 0:
            assert f"${nominal}${stamp['iterations_run']}" in formula
        else:
            assert ",NA())" in formula or ",NA()," in formula, index
    for formula in _block_formulas("distribution", "lower"):
        assert formula.startswith(f'=IF({published}="",NA(),')


def test_34_the_histogram_plots_the_count_against_the_lower_edge() -> None:
    spec = _by_key()["histogram"]
    assert spec["kind"] == "column"
    assert [series["key"] for series in spec["series"]] == ["count"]
    assert spec["categories"]["key"] == "lower"
    # THE PUBLISHED SIMULATION'S OWN STATE. Repointed at the P8-3 pre-Windows
    # correction: the annual line reads NOT PRODUCED whenever the annual step
    # has not run, which says nothing about a histogram whose data is present
    # and current, and the persisted `(last evaluated)` row was proved live to
    # still read CURRENT after a request change.
    assert spec["state_source"] == "simulation_state"
    assert spec["state_source"] != "distribution_state"
    assert spec["source_block"] == "distribution"


# ===========================================================================
# D. THE ANNUAL CASH FLOW
# ===========================================================================
def test_40_the_annual_chart_tracks_the_published_profile_exactly() -> None:
    """THE PERSISTED ANSWER, NOT A NEW ONE. Each bar is the profile cell the
    accepted P8-1 table publishes for that year, guarded and nothing else."""
    display = {c["field"]: c["column"] for c in _shell()["results"]["annual"]["columns"]}
    first = int(_shell()["results"]["annual"]["first_row"])
    for offset, formula in enumerate(_block_formulas("annual", "annual_nominal")):
        assert f"${display['nominal']}${first + offset}" in formula, (offset, formula)
        assert "SUM(" not in formula, "the annual bar is a running total"
    spec = _by_key()["annual_cash_flow"]
    assert spec["kind"] == "column"
    assert [series["key"] for series in spec["series"]] == ["annual_nominal"]
    assert spec["categories"]["key"] == "calendar_year"
    assert spec["state_source"] == "profile_state"


def test_41_the_year_ordering_is_the_published_one() -> None:
    """ROW ORDER IS YEAR ORDER, and it is taken rather than reconstructed: the
    bridge row for year n reads the display row for year n."""
    first = int(_shell()["results"]["annual"]["first_row"])
    display = {c["field"]: c["column"] for c in _shell()["results"]["annual"]["columns"]}
    for offset, formula in enumerate(_block_formulas("annual", "calendar_year")):
        assert f"${display['calendar_year']}${first + offset}" in formula, offset
    for offset, formula in enumerate(_block_formulas("annual", "project_index")):
        assert f"${display['project_index']}${first + offset}" in formula, offset


# ===========================================================================
# E. THE TORNADO
# ===========================================================================
def test_50_the_tornado_selects_a_window_and_ranks_nothing() -> None:
    """PHASE 7 OWNS THE RANKING. The Sensitivity sheet publishes it in rank
    order, so the first N rows ARE the top N - taking them preserves the order
    by construction, and there is nothing here that could re-sort it."""
    sensitivity = _shell()["sensitivity"]
    columns = {c["key"]: c["column"] for c in sensitivity["columns"]}
    first = int(sensitivity["first_row"])
    names = _block_formulas("drivers", "driver_name")
    rhos = _block_formulas("drivers", "rho")
    assert len(names) == len(rhos) == _projection()["bridge"]["drivers"]["row_count"]
    for index, (name, rho) in enumerate(zip(names, rhos)):
        assert f"{sensitivity['sheet']}!${columns['driver_name']}${first + index}" in name
        assert f"{sensitivity['sheet']}!${columns['rho']}${first + index}" in rho
    # NOT ONE SORTING, RANKING OR ABSOLUTE-VALUE CONSTRUCT ANYWHERE.
    for formula in names + rhos:
        for banned in ("SORT", "RANK", "LARGE", "SMALL", "ABS(", "INDEX", "MATCH"):
            assert banned not in formula, f"the tornado re-derives its order: {formula}"


def test_51_the_signed_rho_is_preserved_and_the_abs_column_is_not_plotted() -> None:
    """DIRECTION IS THE POINT OF A TORNADO. Plotting |rho| would draw every bar
    to the right and lose which drivers push the total up and which pull it
    down - and the Sensitivity sheet publishes both columns, so the wrong one is
    one character away."""
    sensitivity = _shell()["sensitivity"]
    columns = {c["key"]: c["column"] for c in sensitivity["columns"]}
    assert "abs_rho" in columns, "the accepted sheet no longer publishes |rho|"
    for formula in _block_formulas("drivers", "rho"):
        assert f"${columns['abs_rho']}$" not in formula, (
            "the tornado plots the absolute value and loses direction")
        assert f"${columns['rho']}$" in formula
    spec = _by_key()["tornado"]
    assert [series["key"] for series in spec["series"]] == ["rho"]
    assert spec["kind"] == "bar", "a tornado is horizontal"
    assert spec["categories"]["key"] == "driver_name"


def test_52_absent_drivers_are_not_drawn_as_zero_length_bars() -> None:
    """FEWER ELIGIBLE DRIVERS THAN N MUST DRAW FEWER BARS. N-k bars of nothing
    sitting on the axis look measured, and zero-variance drivers are excluded by
    the accepted contract precisely so they are not shown as uncorrelated."""
    for formula in _block_formulas("drivers", "driver_name") + _block_formulas("drivers", "rho"):
        assert '="",NA()' in formula.replace(" ", ""), formula
        assert ',""' not in formula, "an absent driver becomes a blank a chart plots as zero"


def test_53_an_unavailable_ranking_is_stated_rather_than_drawn_empty() -> None:
    """A CONVINCING EMPTY TORNADO IS THE WORST OUTCOME. The Sensitivity sheet
    already writes a sentence saying whether its ranking exists and which run it
    belongs to; the bridge mirrors that sentence and the Dashboard shows it."""
    status = _projection()["bridge"]["status"]
    assert status["kind"] == "mirrored"
    keys = {row["key"] for row in status["rows"]}
    assert "sensitivity_availability" in keys
    sensitivity = _shell()["sensitivity"]
    availability = (f"{sensitivity['sheet']}!"
                    f"${sensitivity['columns'][1]['column']}$"
                    f"{sensitivity['availability_row']}")
    availability_row = next(row["row"] for row in status["rows"]
                            if row["key"] == "sensitivity_availability")
    formula = _formula(f"{_shell()['results']['nominal_column']}{availability_row}")
    assert availability in formula, formula
    # MIRRORED, NEVER RE-DERIVED: no arm of that sentence is rebuilt here.
    for banned in ("PUBLISHED", "CURRENT", "IF(AND(", "COUNTIF"):
        assert banned not in formula, f"the bridge re-derives the availability: {banned}"
    # AND THE DASHBOARD SHOWS IT.
    dashboard = json.loads(
        (BUILD / "phase8_dashboard_inspection.json").read_text(encoding="utf-8"))
    mirrored = {(entry["source_block"], entry["source_key"])
                for section in dashboard["sections"] for entry in section["rows"]}
    assert ("chart_status", "sensitivity_availability") in mirrored
    assert _by_key()["tornado"]["state_source"] == "sensitivity_availability"


# ===========================================================================
# F. STATE - NO CHART OWNS ONE
# ===========================================================================
def test_60_every_chart_names_a_state_source_it_does_not_own() -> None:
    """A CHART CARRIES NO STATE LOGIC. It inherits the state its source already
    publishes, and the projection says where a reader finds it."""
    published = {"distribution_state", "profile_state", "profile_px",
                 "sensitivity_availability", "simulation_state"}
    for spec in _by_key().values():
        assert spec["state_source"] in published, spec
        assert spec["no_data_value"] == "NA()"


def test_61_no_chart_bridge_formula_carries_a_state_word_of_its_own() -> None:
    """THE ONE STATE WORD THE BRIDGE MAY SPELL is the NOT PRODUCED the accepted
    annual table already guards on - and it spells it because it reuses that
    exact guard. Every other word would be a second vocabulary."""
    everything = []
    for block in ("annual", "distribution", "drivers"):
        entry = _projection()["bridge"][block]
        for column in entry["columns"]:
            everything += _block_formulas(block, column["key"])
    for formula in everything:
        for word in ("CURRENT", "HISTORICAL", "OTHER Px", "Reconciled",
                     "NOT RECONCILED", "STALE", "INVALID"):
            assert word not in formula, f"a bridge formula spells {word!r}: {formula}"
    assert any("NOT PRODUCED" in formula for formula in everything), (
        "the annual bridge no longer reuses the accepted empty-state guard")


# ===========================================================================
# G. MUTATIONS - each is a way a chart lies convincingly
# ===========================================================================
def _mutated(mutate) -> Path:
    raw = yaml.safe_load(MANIFEST.read_text(encoding="utf-8"))
    mutate(raw["phase6_shell"]["charts"], raw["phase6_shell"])
    path = BUILD / "_phase8_charts_probe.yaml"
    path.write_text(yaml.safe_dump(raw, sort_keys=False), encoding="utf-8")
    return path


@pytest.mark.parametrize("name,mutate,expected", [
    # THE WRONG SOURCE SHEET - the bridge no longer lives where the charts read.
    ("the bridge moves off Results",
     lambda charts, shell: charts.__setitem__("bridge_sheet", "Sensitivity"),
     "the bridge lives on"),
    # A SERIES NOTHING PUBLISHES - the P7-4 stale-address shape.
    ("a chart plots a series the bridge does not publish",
     lambda charts, shell: charts["charts"][0]["series"].__setitem__(
         0, {"key": "expected_exposure", "name": "Expected"}),
     "which the annual block does not publish"),
    # THE CATEGORY TAKEN FROM A BLOCK THAT HAS NO SUCH COLUMN.
    ("a chart takes categories from a column that does not exist",
     lambda charts, shell: charts["charts"][1].__setitem__("categories", "calendar_year"),
     "which the distribution block does not publish"),
    # A THIRD DIMENSION.
    ("a chart becomes three-dimensional",
     lambda charts, shell: charts["charts"][0].__setitem__("kind", "bar3D"),
     "is a 'bar3D' chart"),
    # THE BRIDGE GROWN UP INTO THE ACCEPTED P8-1 SURFACE.
    ("the bridge starts inside the accepted Results layout",
     lambda charts, shell: charts["bridge"].__setitem__("heading_row", 100),
     "inside the accepted Results layout"),
    # TWO BRIDGE BLOCKS OVERLAID - the silent one.
    ("two bridge blocks write the same row",
     lambda charts, shell: charts["bridge"]["drivers"].__setitem__(
         "first_row", charts["bridge"]["distribution"]["first_row"]),
     "both write row"),
    # A CHART ANCHORED OVER THE EXECUTIVE SUMMARY.
    ("a chart is anchored above the reserved region",
     lambda charts, shell: charts["charts"][0].__setitem__("anchor", "B20"),
     "outside the reserved region"),
    # A STATE SOURCE NOTHING PUBLISHES.
    ("a chart names a state source nothing publishes",
     lambda charts, shell: charts["charts"][0].__setitem__("state_source", "model_status"),
     "which nothing publishes"),
])
def test_70_the_manifest_refuses_each_way_a_chart_lies(
        name: str, mutate, expected: str) -> None:
    """EACH MUTATION APPLIED TO A COPY OF THE MANIFEST and refused by the loader,
    naming what is wrong. Nothing on disk changes."""
    path = _mutated(mutate)
    try:
        with pytest.raises(SpecError, match=re.escape(expected)):
            load_spec(path)
    finally:
        path.unlink(missing_ok=True)


@pytest.mark.parametrize("name,mutate", [
    # THE TWO LADDERS, ONE SHEET OUT: a chart plotting its own axis.
    ("a chart plots its category column as a series",
     lambda inspection: inspection["charts"][0]["series"].__setitem__(
         0, dict(inspection["charts"][0]["categories"], name="Category"))),
    # A RANGE THAT LOST ITS SHEET - it would resolve against the Dashboard,
    # which holds none of this data, and plot nothing while looking fine.
    ("a series range loses its sheet",
     lambda inspection: inspection["charts"][0]["series"][0].__setitem__(
         "range", "$H$279:$H$478")),
    # A LEGEND THAT DISAGREES WITH THE SERIES COUNT.
    ("a single-series chart declares a legend",
     lambda inspection: inspection["charts"][1].__setitem__("legend", True)),
    # TWO CHARTS IN ONE PLACE.
    ("two charts share an anchor",
     lambda inspection: inspection["charts"][1].__setitem__(
         "anchor", inspection["charts"][0]["anchor"])),
    # THE SOURCE THE TORNADO MIRRORS, SIMPLY GONE - the second Windows run's
    # crash, seen from the side that can refuse it before Excel is opened.
    ("the sensitivity source is missing",
     lambda inspection: inspection.pop("sensitivity_source")),
    ("the sensitivity source loses its first row",
     lambda inspection: inspection["sensitivity_source"].pop("first_row")),
    ("the sensitivity source loses its columns",
     lambda inspection: inspection["sensitivity_source"].pop("columns")),
    # THE MAGNITUDE MIRRORED INSTEAD OF THE SIGNED CORRELATION. Every driver
    # would still plot; every negative one would lose its direction.
    ("the tornado mirrors the absolute magnitude",
     lambda inspection: inspection["sensitivity_source"]["columns"][1].__setitem__(
         "key", "abs_rho")),
    # A SOURCE FIELD THE BRIDGE DOES NOT PLOT.
    ("the source mirrors a field the bridge does not plot",
     lambda inspection: inspection["sensitivity_source"]["columns"][0].__setitem__(
         "key", "driver_id")),
    # THE TWO SIDES IN DIFFERENT ORDERS - a positional comparison would then
    # check the name against the rho.
    ("the source and the bridge disagree about order",
     lambda inspection: inspection["sensitivity_source"]["columns"].reverse()),
    # A ROW NUMBER WHERE A COLUMN LETTER BELONGS.
    ("a source column stops being a column",
     lambda inspection: inspection["sensitivity_source"]["columns"][0].__setitem__(
         "column", "13")),
    # THE RANKING PRODUCED BY A CELL RATHER THAN BY AN OPERATOR. A worksheet
    # function here would mean the tornado could rerun the analysis in order to
    # draw itself, which is the second analytical authority this layer refuses.
    ("the sensitivity endpoint becomes a worksheet function",
     lambda inspection: inspection.__setitem__(
         "sensitivity_endpoint", "PCCM_SensitivityRanking")),
])
def test_71_the_projection_validator_refuses_each_incoherent_chart(
        name: str, mutate) -> None:
    """THE LOADER CANNOT SEE THESE - they are structurally legal manifests whose
    PROJECTION is incoherent - so the projection's own validator must."""
    structure = load_structure_contract(SPEC / "structure_contract.yaml")
    inspection = build_phase8_charts_inspection(
        load_spec(MANIFEST), structure.limits.max_generated_year_columns,
        _sim_contract())
    mutate(inspection)
    with pytest.raises(ValueError):
        validate_phase8_charts_inspection(inspection)


def test_72_the_unmutated_manifest_passes_both_gates() -> None:
    """SO THE TWENTY REFUSALS ABOVE ARE REFUSALS OF THE MUTATION, not of the
    fixture."""
    structure = load_structure_contract(SPEC / "structure_contract.yaml")
    inspection = build_phase8_charts_inspection(
        load_spec(MANIFEST), structure.limits.max_generated_year_columns,
        _sim_contract())
    validate_phase8_charts_inspection(inspection)
    assert inspection == _projection(), (
        "the committed chart projection is not what the manifest now produces")


def test_72a_the_sensitivity_endpoint_is_projected_from_the_manifest() -> None:
    """THE TORNADO'S PRODUCER IS DECLARED, NOT DISCOVERED. A Windows runner that
    had to type the endpoint would be a second declaration of something the
    manifest owns, and P7-4 is what that costs when the first one moves.

    NOTHING CALLS IT AUTOMATICALLY. Naming it here is not a licence for a chart,
    an adapter or a recalculation to produce a ranking; sensitivity is produced
    when an operator asks, and the tornado says so when nobody has."""
    spec = load_spec(MANIFEST)
    declared = spec.phase6_shell["charts"]["sensitivity_endpoint"]
    structure = load_structure_contract(SPEC / "structure_contract.yaml")
    inspection = build_phase8_charts_inspection(
        spec, structure.limits.max_generated_year_columns,
        _sim_contract())
    assert inspection["sensitivity_endpoint"] == declared
    assert declared == "PCCM_RunSensitivity"
    # AND IT IS AN ENDPOINT THE COMMAND SURFACE ACTUALLY PUBLISHES - a Sub, so
    # no cell can reach it.
    source = (SRC / "modSimPostReport.bas").read_text(encoding="utf-8")
    assert f"Public Sub {declared}()" in source, (
        f"{declared} is not a published command endpoint")


# ===========================================================================
# H. THE PRE-WINDOWS STATE CORRECTION
# ===========================================================================
# WHAT WAS WRONG, AND IT WAS WRONG IN TWO DIRECTIONS.
#
# THE HISTOGRAM plots the published SIMULATION distribution and was qualified by
# the ANNUAL distribution state. That state consumes the live simulation status,
# so it could never say CURRENT while the simulation was stale - safe in the
# over-claiming direction - but it reads NOT PRODUCED whenever the annual step
# has not run, which says nothing at all about a histogram whose data is present
# and current.
#
# THE TORNADO was worse. Its only qualifier compared two PERSISTED records - the
# sensitivity block's stored fingerprint against the published run's stored
# fingerprint - and is therefore structurally blind to a model that has moved
# since. After a request change with nothing rerun it still read
# "CURRENT for run N" while the model had moved past run N. That is
# over-claiming, and no amount of reading it more carefully fixes it.
#
# WHAT NEITHER MAY EVER USE is Results row 23, the persisted
# `Simulation Status (last evaluated)`. P8-1 proved in live Excel that it can
# still read CURRENT after a request change.

LIVE_ADAPTER = "PCCM_ResultsSimulationState"
PURE_OWNER = "SimReportDerivedStatus"
WRITING_PATH = "PCCM_SimulationStatus"


def _adapter_source() -> str:
    return (SRC / "modResultsState.bas").read_text(encoding="utf-8")


def _adapter_body(name: str) -> str:
    source = _adapter_source()
    start = source.index(f"Public Function {name}")
    return source[start:source.index("End Function", start)]


def _vba_procedures() -> dict[tuple[str, str], str]:
    if "vba" not in _CACHE:
        from pccm_builder.vba_source import load_modules

        out: dict[tuple[str, str], str] = {}
        for module in load_modules([SRC]):
            lines = module.code.splitlines()
            index = 0
            while index < len(lines):
                head = re.match(
                    r"\s*(?:Public |Private |Friend )?(?:Static )?"
                    r"(?:Function|Sub|Property (?:Get|Let|Set))\s+([A-Za-z_]\w*)",
                    lines[index])
                if head:
                    end = index + 1
                    while end < len(lines) and not re.match(
                            r"\s*End (?:Function|Sub|Property)\b", lines[end]):
                        end += 1
                    out[(module.name, head.group(1))] = "\n".join(lines[index + 1:end])
                    index = end
                index += 1
        _CACHE["vba"] = out
    return _CACHE["vba"]


def _reachable(root: tuple[str, str]) -> set[tuple[str, str]]:
    procedures = _vba_procedures()
    seen: set[tuple[str, str]] = set()
    stack = [root]
    while stack:
        current = stack.pop()
        if current in seen:
            continue
        seen.add(current)
        body = procedures.get(current, "")
        found: set[tuple[str, str]] = set()
        for match in re.finditer(r"\b(mod\w+)\.([A-Za-z_]\w*)", body):
            found.add((match.group(1), match.group(2)))
        for match in re.finditer(r"\b([A-Za-z_]\w*)\b", body):
            if (current[0], match.group(1)) in procedures:
                found.add((current[0], match.group(1)))
        for callee in found:
            if callee in procedures and callee not in seen:
                stack.append(callee)
    return seen


_WORKBOOK_WRITE = re.compile(
    r"""(?:\.Value2|\.Value|\.Formula\w*|\.NumberFormat|\.Text)\s*=(?!=)"""
    r"""|\.ClearContents\b|\.EntireRow\.Delete\b|ListRows\.Add\b""")
_UDF_HOSTILE = re.compile(
    r"\.Select\b|\.Activate\b|MsgBox|Application\.Run\b|\.Calculate\b"
    r"|Application\.(?:ScreenUpdating|EnableEvents|Calculation|DisplayAlerts)\s*="
    r"|ActiveSheet|Selection\b|DoEvents|\.SaveAs\b")


def test_80_the_live_adapter_is_a_thin_volatile_delegation() -> None:
    """THE SAME SHAPE AS THE FOUR P8-1 ADAPTERS, and nothing more: volatile
    because a zero-argument function has no input Excel can watch, loud on
    error because a plausible wrong state word is worse than an error, and one
    statement because a presentation layer owns no semantic."""
    body = _adapter_body(LIVE_ADAPTER)
    assert "Application.Volatile True" in body
    assert "On Error GoTo Unavailable" in body
    assert "CVErr(xlErrValue)" in body
    assert "Resume Next" not in body, "the adapter swallows the error"
    statements = [line.strip() for line in body.splitlines()
                  if line.strip() and not line.strip().startswith("'")
                  and not line.strip().startswith(("Public", "Exit", "Unavailable"))]
    assert statements == [
        "On Error GoTo Unavailable",
        "Application.Volatile True",
        f"{LIVE_ADAPTER} = modSimReport.{PURE_OWNER}()",
        f"{LIVE_ADAPTER} = CVErr(xlErrValue)",
    ], statements


def test_81_it_delegates_to_the_accepted_pure_owner_and_never_to_the_writer() -> None:
    """THE PURE HALF, BY NAME. P8-1 split `DeriveSimStatus` out from the entry
    point that also persists it, precisely so a worksheet cell could ask without
    asking for a rewrite. This is that split being used."""
    body = _adapter_body(LIVE_ADAPTER)
    assert f"modSimReport.{PURE_OWNER}()" in body
    assert WRITING_PATH not in body, "the adapter reaches the writing status path"
    reachable = _reachable(("modResultsState", LIVE_ADAPTER))
    assert ("modSimReport", PURE_OWNER) in reachable
    assert ("modSimReport", "DeriveSimStatus") in reachable
    for forbidden in (("modSimReport", WRITING_PATH),
                      ("modSimReport", "WriteStatusBlock"),
                      ("modCalcReport", "PCCM_CalculationStatus"),
                      ("modCalcReport", "WriteStatusBlock")):
        assert forbidden not in reachable, (
            f"{LIVE_ADAPTER} reaches {forbidden[0]}.{forbidden[1]}")


def test_82_the_whole_call_graph_of_the_live_adapter_is_read_only() -> None:
    """TAKEN OVER THE TRANSITIVE CLOSURE, not over the adapter's four lines. A
    cell may not change the book, so nothing a cell can call may either -
    however many hops away."""
    procedures = _vba_procedures()
    reachable = _reachable(("modResultsState", LIVE_ADAPTER))
    assert len(reachable) > 100, len(reachable)
    offenders = [
        f"{module}.{name}: {line.strip()[:70]}"
        for module, name in sorted(reachable)
        for line in procedures.get((module, name), "").splitlines()
        if _WORKBOOK_WRITE.search(line) or _UDF_HOSTILE.search(line)]
    assert not offenders, "\n  ".join(offenders)


def test_83_no_state_arm_is_duplicated_in_the_adapter_module() -> None:
    """A RULE IS A DECISION, NOT A DELEGATION. The module may call the owner of
    a semantic; it may never branch on one."""
    code = "\n".join(line for line in _adapter_source().splitlines()
                     if not line.lstrip().startswith("'"))
    for rule in ("SIM_STATE_", "SIM_ANNUAL_STATE_", "CURRENT", "STALE", "INVALID",
                 "HISTORICAL", "NOT PRODUCED", "StrComp", "If ", "Select Case",
                 "DeriveSimStatus", "WriteStatusBlock"):
        assert rule not in code, f"the adapter module decides a state: {rule}"
    assert code.count(f"modSimReport.{PURE_OWNER}()") == 1


def test_84_the_histogram_is_qualified_by_the_live_simulation_state() -> None:
    """NOT THE ANNUAL STATE, AND NOT THE PERSISTED ROW. Both were wrong for
    different reasons, and both are asserted against by name."""
    spec = _by_key()["histogram"]
    assert spec["state_source"] == "simulation_state"
    assert spec["state_source"] not in ("distribution_state", "profile_state")
    status = _projection()["bridge"]["status"]
    row = next(entry["row"] for entry in status["rows"]
               if entry["key"] == "simulation_state")
    formula = _formula(f"{_shell()['results']['nominal_column']}{row}")
    assert formula == f"={LIVE_ADAPTER}()", formula
    # AND IT IS NOT THE `(last evaluated)` ROW. P8-1 proved live that that one
    # can still read CURRENT after a request change.
    stamp = {f["key"]: f["row"] for f in _results_projection()["run_stamp"]["fields"]}
    persisted = f"${_results_projection()['columns']['nominal']}${stamp['simulation_status']}"
    assert persisted not in formula, "the histogram borrows the persisted status row"


def test_85_request_drift_and_an_invalid_model_cannot_leave_it_current() -> None:
    """THE MECHANISM, TRACED TO THE SOURCE rather than asserted. The adapter
    delegates to `DeriveSimStatus`, which recomputes the CURRENT request
    fingerprint from live inputs and compares it to the published one - so a
    changed request yields STALE and an unformable model yields INVALID, with no
    endpoint invoked and nothing persisted."""
    report = (SRC / "modSimReport.bas").read_text(encoding="utf-8")
    derive = report[report.index("Private Function DeriveSimStatus"):]
    derive = derive[:derive.index("\nEnd Function")]
    assert "CurrentRequestFingerprint(fingerprint, detail)" in derive, (
        "the derivation no longer recomputes the current request")
    assert "SIM_STATE_STALE" in derive and "SIM_STATE_INVALID" in derive
    assert "SIM_STATE_CURRENT" in derive
    # THE WORD IS PASSED THROUGH UNTRANSLATED. Folding STALE and INVALID into
    # HISTORICAL would invent a Phase-8 vocabulary for a caption's convenience.
    body = _adapter_body(LIVE_ADAPTER)
    for word in ("HISTORICAL", "STALE", "INVALID", "CURRENT"):
        assert word not in body, f"the adapter translates the state word {word!r}"
    # AND THE RECOMPUTATION IS SIDE-EFFECT FREE.
    current = report[report.index("Private Function CurrentRequestFingerprint"):]
    current = current[:current.index("\nEnd Function")]
    assert "SIDE-EFFECT FREE" in current


def test_86_the_tornado_carries_two_conditions_that_do_not_collapse() -> None:
    """TWO QUESTIONS, TWO ANSWERS. "Does this ranked table belong to the
    published run?" is a comparison of two persisted records and is blind to a
    model that has moved since; "does that run still match the model?" is the
    live one. Either alone said CURRENT after the request drifted."""
    spec = _by_key()["tornado"]
    assert spec["state_source"] == "sensitivity_availability"
    assert spec["also_qualified_by"] == "simulation_state"
    assert spec["state_source"] != spec["also_qualified_by"]
    # AND THE AVAILABILITY SENTENCE REALLY IS BLIND ON ITS OWN - which is why
    # the second condition exists. It calls no UDF and reads only persisted
    # cells, so nothing in it can notice a model change.
    availability = _shell()["sensitivity"]["availability_formula"]
    assert "PCCM_" not in availability, (
        "the availability sentence now reaches a live evaluator; the second "
        "condition may have become redundant and this control should be revisited")
    assert "_SimData!" in availability
    # THE OTHER THREE CHARTS DECLARE NO SECOND CONDITION, so this is not a
    # blanket field nobody reads.
    for key in ("s_curve", "histogram", "annual_cash_flow"):
        assert _by_key()[key]["also_qualified_by"] is None, key


def test_87_both_qualifications_are_visible_with_the_charts() -> None:
    """THE PREVIOUS ARRANGEMENT PUT THE STATE FORTY ROWS ABOVE THE PLOTS, which
    is no qualification at all for a reader who has scrolled. Two things fix it
    and neither creates a second authority: the executive status block is frozen
    on screen, and the two chart-specific lines sit inside the region,
    immediately above the first plot."""
    dashboard = json.loads(
        (BUILD / "phase8_dashboard_inspection.json").read_text(encoding="utf-8"))
    region = dashboard["chart_region"]
    status = next(section for section in dashboard["sections"]
                  if section["key"] == "chart_status")
    keys = {entry["key"] for entry in status["rows"]}
    assert keys == {"simulation_state", "sensitivity_availability"}, keys
    # INSIDE THE REGION, ABOVE EVERY PLOT.
    for entry in status["rows"]:
        assert int(region["heading_row"]) < entry["row"] < int(region["first_row"]), entry
    first_chart = min(int(re.sub(r"[A-Z]", "", chart["anchor"]))
                      for chart in _projection()["charts"])
    assert max(entry["row"] for entry in status["rows"]) < first_chart
    assert first_chart - max(entry["row"] for entry in status["rows"]) <= 4, (
        "the chart status is no longer adjacent to the plots it qualifies")
    # AND THE EXECUTIVE STATUS BLOCK IS FROZEN ON SCREEN.
    sheet = _workbook()[dashboard["sheet"]]
    assert sheet.freeze_panes, "the Dashboard does not freeze its status block"
    frozen_below = int(re.sub(r"[A-Z]", "", sheet.freeze_panes))
    summary = next(section for section in dashboard["sections"]
                   if section["key"] == "status")
    assert max(entry["row"] for entry in summary["rows"]) < frozen_below, (
        "the frozen pane does not keep the whole Result Status block on screen")


def test_88_the_chart_status_reaches_the_dashboard_only_through_results() -> None:
    """ONE SURFACE. The live state is a Results cell and the availability is a
    Results cell; the Dashboard mirrors both from there. A Dashboard formula
    reading Sensitivity - or `_SimData` - would be a second surface for the
    sheet P8-2 established reads one."""
    sheet = _workbook()[_projection()["chart_sheet"]]
    for row in sheet.iter_rows():
        for cell in row:
            if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                continue
            for forbidden in ("_SimData", "_Calc", "Sensitivity!"):
                assert forbidden not in cell.value, (
                    f"Dashboard!{cell.coordinate} reads {forbidden}: {cell.value!r}")
            assert cell.value.count("Results!") == 2, cell.coordinate


def test_89_the_p8_1_and_p8_2_accepted_geometry_still_holds() -> None:
    """THE CORRECTION ADDS; IT DOES NOT MOVE ANYTHING ACCEPTED."""
    p81 = yaml.safe_load(_git("show", f"{P81_ACCEPTANCE}:pccm/spec/workbook.yaml"))
    assert p81["phase6_shell"]["results"] == _shell()["results"]
    p82 = yaml.safe_load(_git("show", f"{P82_ACCEPTANCE}:pccm/spec/workbook.yaml"))
    accepted = p82["phase6_shell"]["dashboard"]
    current = _shell()["dashboard"]
    for key in ("sheet", "source_sheet", "label_column", "nominal_column",
                "pv_column", "mirror_formula", "number_formats"):
        assert accepted[key] == current[key], key
    assert current["sections"][:len(accepted["sections"])] == accepted["sections"], (
        "an accepted P8-2 summary section changed")


@pytest.mark.parametrize("name,mutate", [
    # THE HISTOGRAM PUT BACK ON THE ANNUAL STATE.
    ("the histogram is repointed at the annual state",
     lambda charts, shell: charts["charts"][1].__setitem__(
         "state_source", "distribution_state")),
    # THE HISTOGRAM PUT ON THE PERSISTED `(last evaluated)` ROW.
    ("the histogram is repointed at the persisted status row",
     lambda charts, shell: charts["charts"][1].__setitem__(
         "state_source", "simulation_status")),
    # THE TORNADO LOSING ITS LIVE CONDITION.
    ("the tornado stops asking whether the run matches the model",
     lambda charts, shell: charts["charts"][3].pop("also_qualified_by")),
    # THE TWO CONDITIONS COLLAPSED INTO ONE.
    ("the tornado names one condition twice",
     lambda charts, shell: charts["charts"][3].__setitem__(
         "also_qualified_by", "sensitivity_availability")),
    # THE LIVE STATE LINE REMOVED FROM THE BRIDGE ENTIRELY.
    ("the bridge stops publishing the live simulation state",
     lambda charts, shell: charts["bridge"]["status"].__setitem__(
         "rows", [row for row in charts["bridge"]["status"]["rows"]
                  if row["key"] != "simulation_state"])),
])
def test_90_each_way_of_losing_the_correction_is_refused(name: str, mutate) -> None:
    """EACH MUTATION APPLIED TO A COPY OF THE MANIFEST. The loader refuses the
    ones it can see; the rest are caught by the state-source rules above, re-run
    here against the mutated projection."""
    path = _mutated(mutate)
    try:
        try:
            spec = load_spec(path)
        except SpecError:
            return  # refused at the gate, which is the strongest outcome
        structure = load_structure_contract(SPEC / "structure_contract.yaml")
        inspection = build_phase8_charts_inspection(
            spec, structure.limits.max_generated_year_columns,
            _sim_contract())
        histogram = next(c for c in inspection["charts"] if c["key"] == "histogram")
        tornado = next(c for c in inspection["charts"] if c["key"] == "tornado")
        broken = (
            histogram["state_source"] != "simulation_state"
            or tornado["also_qualified_by"] != "simulation_state"
            or tornado["also_qualified_by"] == tornado["state_source"]
            or "simulation_state" not in {row["key"]
                                          for row in inspection["bridge"]["status"]["rows"]})
        assert broken, f"'{name}' left the correction intact"
    finally:
        path.unlink(missing_ok=True)


# THE RULES THE ADAPTER MUST SATISFY, AS FUNCTIONS OVER ITS SOURCE, so the
# mutations below are refused by the SAME checks that pass on the real module
# rather than by a restatement of them.
def _adapter_is_thin_and_volatile(source: str) -> None:
    start = source.index(f"Public Function {LIVE_ADAPTER}")
    body = source[start:source.index("End Function", start)]
    assert "Application.Volatile True" in body, (
        "the adapter is not volatile; a zero-argument cell would answer once")
    assert "On Error GoTo Unavailable" in body and "CVErr(xlErrValue)" in body
    statements = [line.strip() for line in body.splitlines()
                  if line.strip() and not line.strip().startswith("'")
                  and not line.strip().startswith(("Public", "Exit", "Unavailable"))]
    assert statements == [
        "On Error GoTo Unavailable",
        "Application.Volatile True",
        f"{LIVE_ADAPTER} = modSimReport.{PURE_OWNER}()",
        f"{LIVE_ADAPTER} = CVErr(xlErrValue)",
    ], statements


def _adapter_never_writes(source: str) -> None:
    code = "\n".join(line for line in source.splitlines()
                     if not line.lstrip().startswith("'"))
    assert WRITING_PATH not in code, "the adapter reaches the writing status path"
    assert "WriteStatusBlock" not in code


def _adapter_owns_no_rule(source: str) -> None:
    code = "\n".join(line for line in source.splitlines()
                     if not line.lstrip().startswith("'"))
    for rule in ("SIM_STATE_", "SIM_ANNUAL_STATE_", "CURRENT", "STALE", "INVALID",
                 "HISTORICAL", "NOT PRODUCED", "StrComp", "If ", "Select Case",
                 "DeriveSimStatus"):
        assert rule not in code, f"the adapter module decides a state: {rule}"


ADAPTER_RULES = (_adapter_is_thin_and_volatile, _adapter_never_writes,
                 _adapter_owns_no_rule)


@pytest.mark.parametrize("name,mutate", [
    # THE ADAPTER CALLING THE WRITING PATH - the P8-1 defect, reintroduced.
    ("the adapter calls the writing status path",
     lambda code: code.replace(f"modSimReport.{PURE_OWNER}()",
                               f"modSimReport.{WRITING_PATH}()")),
    # STATE LOGIC DUPLICATED INTO THE ADAPTER MODULE.
    ("a state arm is duplicated in the adapter",
     lambda code: code.replace(
         f"    {LIVE_ADAPTER} = modSimReport.{PURE_OWNER}()",
         f"    If modSimReport.{PURE_OWNER}() = SIM_STATE_STALE Then\n"
         f"        {LIVE_ADAPTER} = \"HISTORICAL\"\n"
         f"    Else\n"
         f"        {LIVE_ADAPTER} = modSimReport.{PURE_OWNER}()\n"
         f"    End If")),
    # STALE AND INVALID FOLDED INTO HISTORICAL FOR A TIDIER CAPTION.
    ("the adapter translates the state word",
     lambda code: code.replace(
         f"    {LIVE_ADAPTER} = modSimReport.{PURE_OWNER}()",
         f"    {LIVE_ADAPTER} = \"HISTORICAL\"")),
    # THE VOLATILITY REMOVED - the cell would answer once and never again.
    ("the adapter stops being volatile",
     lambda code: code.replace(
         f"    Application.Volatile True\n    {LIVE_ADAPTER} =",
         f"    {LIVE_ADAPTER} =")),
    # THE ERROR SWALLOWED - a plausible wrong word instead of a visible failure.
    ("the adapter swallows its error",
     lambda code: code.replace(
         f"    {LIVE_ADAPTER} = CVErr(xlErrValue)", f"    {LIVE_ADAPTER} = \"\"")),
])
def test_91_each_production_mutation_of_the_adapter_is_refused(
        name: str, mutate) -> None:
    """APPLIED TO A COPY OF THE MODULE IN MEMORY and run against the same three
    rules the real module satisfies. Nothing on disk changes."""
    mutated = mutate(_adapter_source())
    assert mutated != _adapter_source(), f"'{name}' changed nothing"
    refused = []
    for rule in ADAPTER_RULES:
        try:
            rule(mutated)
        except (AssertionError, ValueError) as failure:
            refused.append(f"{rule.__name__}: {failure}")
    assert refused, f"'{name}' survived every adapter rule"


def test_92_the_adapter_rules_pass_on_the_real_module() -> None:
    """SO THE FIVE REFUSALS ABOVE ARE REFUSALS OF THE MUTATION, not of the
    fixture."""
    for rule in ADAPTER_RULES:
        rule(_adapter_source())


# ===========================================================================
# I. THE RESTATED OWNERSHIP CONTROLS ARE NOT VACUOUS
# ===========================================================================
# THE TWO RESTATEMENTS THIS SECTION GUARDS. A control that changed from
# "nothing may change" to "declared changes may" is only worth what its refusals
# are worth, so each refusal is exercised against a deliberately broken copy of
# the evidence rather than described.


class _FakeGit:
    """A git that answers with the diff a mutation wants tested. Nothing on
    disk changes and no repository is touched."""

    def __init__(self, name_status: str, per_path: dict[str, str] | None = None) -> None:
        self._name_status = name_status
        self._per_path = per_path or {}

    def __call__(self, *args: str) -> str:
        if "--name-status" in args:
            return self._name_status
        return self._per_path.get(args[-1], "")


@pytest.mark.parametrize("name,git,expected", [
    # AN UNDECLARED MODULE CHANGED - the thing this control has always been for.
    ("an undeclared production module changed",
     _FakeGit("M\tpccm/src/vba/modSimEngine.bas\n"
              "M\tpccm/src/vba/modResultsState.bas\n"),
     "without being declared"),
    # A MODULE DELETED - no declaration may ever permit this.
    ("a production module was deleted",
     _FakeGit("D\tpccm/src/vba/modSimStats.bas\n"
              "M\tpccm/src/vba/modResultsState.bas\n"),
     "was removed"),
    # A DECLARED FILE EDITED RATHER THAN EXTENDED. A declaration buys the right
    # to ADD to a module, never to rewrite an accepted procedure in it.
    ("a declared file lost a line",
     _FakeGit("M\tpccm/src/vba/modResultsState.bas\n",
              {"pccm/src/vba/modResultsState.bas":
               "--- a/x\n+++ b/x\n-    Application.Volatile True\n+    Nothing\n"}),
     "does not name"),
    # A DECLARED FILE LOSING A LINE ITS DECLARATION DOES NOT NAME. Naming one
    # removal does not open the file: everything else is still an edit to an
    # accepted procedure.
    ("a file declared for one removal loses a different line",
     _FakeGit("M\tpccm/src/vba/modSimPostReport.bas\n"
              "M\tpccm/src/vba/modResultsState.bas\n",
              {"pccm/src/vba/modSimPostReport.bas":
               "--- a/x\n+++ b/x\n"
               "-    column = COL_RISK_REGISTER_DESCRIPTION\n"
               "+    column = COL_RISK_REGISTER_RISK_NAME\n"
               "-        block(row, SIM_SENSITIVITY_OFFSET_RANK + 1) = rank\n"}),
     "does not name"),
    # A NAMED REMOVAL HELD IN RESERVE. A declaration that permits a removal
    # which never happened is a licence waiting for the next change to use.
    ("a declared removal never happened",
     _FakeGit("M\tpccm/src/vba/modSimPostReport.bas\n"
              "M\tpccm/src/vba/modResultsState.bas\n",
              {"pccm/src/vba/modSimPostReport.bas":
               "--- a/x\n+++ b/x\n+    ' a comment\n"}),
     "declares it removes"),
    # A DECLARATION FOR A FILE NOBODY TOUCHED - decoration the next real change
    # would hide beside.
    ("the declared correction changed nothing",
     _FakeGit(""),
     "declared corrections that changed nothing"),
])
def test_93_the_declared_production_rule_refuses_each_undeclared_shape(
        name: str, git, expected: str) -> None:
    with pytest.raises(AssertionError, match=re.escape(expected)):
        _declared_production_changes(git, P82_ACCEPTANCE)


def test_94_the_declared_production_rule_passes_on_the_real_repository() -> None:
    """SO THE SIX REFUSALS ABOVE ARE REFUSALS OF THE MUTATION. And each declared
    correction really is what it says it is."""
    _declared_production_changes(_git, P81_ACCEPTANCE)
    _declared_production_changes(_git, P82_ACCEPTANCE)
    # THE EXACT SET, AND IT GREW BY ONE AT P9-2. Naming them keeps this as
    # strict as it was: a FOURTH declaration still fails here, and so does a
    # removal or a rename of any of these three.
    assert set(DECLARED_PRODUCTION_CORRECTIONS) == {
        "pccm/src/vba/modCalcReport.bas",
        "pccm/src/vba/modResultsState.bas",
        "pccm/src/vba/modSimPostReport.bas"}, sorted(DECLARED_PRODUCTION_CORRECTIONS)
    reason, removals = DECLARED_PRODUCTION_CORRECTIONS[
        "pccm/src/vba/modResultsState.bas"]
    assert LIVE_ADAPTER in reason and PURE_OWNER in reason, reason
    assert removals == (), "the adapter addition is additive and stays so"
    # THE P9-2 SPLIT, ON THE SAME TERMS: additive only, and the declaration says
    # what it exposes, what it delegates to, and what it leaves alone.
    reason, removals = DECLARED_PRODUCTION_CORRECTIONS[
        "pccm/src/vba/modCalcReport.bas"]
    assert removals == (), "the pure-derivation exposure is additive and stays so"
    assert "CalcReportDerivedStatus" in reason, reason
    assert "DeriveStatus" in reason and "PrepareCurrentCalculation" in reason, reason
    assert "PCCM_CalculationStatus" in reason and "WriteStatusBlock" in reason, (
        "the declaration does not say which entry point still persists")
    # THE LABEL CORRECTION REMOVES EXACTLY ONE LINE, and names which.
    reason, removals = DECLARED_PRODUCTION_CORRECTIONS[
        "pccm/src/vba/modSimPostReport.bas"]
    assert removals == ("    column = COL_RISK_REGISTER_DESCRIPTION",), removals
    assert "COL_RISK_REGISTER_RISK_NAME" in reason
    assert "required: true" in reason and "required: false" in reason, (
        "the declaration does not say what made the old column the wrong one")


def test_94a_the_label_correction_is_the_only_thing_that_module_lost() -> None:
    """THE NARROW EXTENSION, CHECKED AGAINST THE REPOSITORY. Whatever else the
    module gained, the single line it lost is the one the declaration names."""
    path = "pccm/src/vba/modSimPostReport.bas"
    removed = [line[1:] for line in
               _git("diff", f"{P81_ACCEPTANCE}..HEAD", "--", path).splitlines()
               if line.startswith("-") and not line.startswith("---")]
    assert removed == ["    column = COL_RISK_REGISTER_DESCRIPTION"], removed
    # AND THE MEASURE AND ORDERING LINES ARE STILL THERE, unmoved.
    source = (SRC / "modSimPostReport.bas").read_text(encoding="utf-8")
    for kept in ("block(row, SIM_SENSITIVITY_OFFSET_RHO + 1) = record.Rho",
                 "block(row, SIM_SENSITIVITY_OFFSET_ABS_RHO + 1) = record.AbsRho",
                 "block(row, SIM_SENSITIVITY_OFFSET_RANK + 1) = rank",
                 "DirectionOf(record.Rho)",
                 "column = COL_COST_LINES_DESCRIPTION"):
        assert kept in source, f"the label correction disturbed {kept!r}"


# ===========================================================================
# J. THE PROCEDURE-LEVEL modSimReport DETECTOR IS NOT VACUOUS
# ===========================================================================
def _modules_with(name: str, extra_code: str):
    """The real module set with one module's code extended in memory."""
    from pccm_builder.vba_source import VbaModule, load_modules

    out = []
    for module in load_modules([SRC]):
        if module.name == name:
            out.append(VbaModule(name=module.name, path=module.path,
                                 raw=module.raw + "\n" + extra_code))
        else:
            out.append(module)
    return out


@pytest.mark.parametrize("name,module,injected,rule", [
    # A FINGERPRINT BUILT INSIDE THE PRESENTATION ADAPTER.
    ("the adapter constructs a fingerprint", "modResultsState",
     "Public Function Probe() As String\n"
     "    Probe = modSimFingerprint.SimFpBuildRequestFingerprint()\n"
     "End Function", "both"),
    # A FINGERPRINT READ INSIDE THE PRESENTATION ADAPTER.
    ("the adapter reads a stored fingerprint", "modResultsState",
     "Public Function Probe() As String\n"
     "    Probe = modSimReport.PCCM_SimulationRequestFingerprint()\n"
     "End Function", "both"),
    # THE PERMITTED DEPENDENCY BROADENED BEYOND THE PURE EVALUATOR.
    ("the adapter calls a second modSimReport procedure", "modResultsState",
     "Public Function Probe() As String\n"
     "    Probe = modSimReport.PCCM_SimulationStatus()\n"
     "End Function", "ownership"),
    # THE RUN ENDPOINT ESCAPING ITS OWNER.
    ("a module acquires the run endpoint", "modSimAnnualStore",
     "Public Function Probe() As String\n"
     "    Probe = PCCM_RunSimulation()\n"
     "End Function", "ownership"),
    # AN UNDECLARED MODULE REACHING THE OWNER AT ALL.
    ("an undeclared module names the owner", "modSimStats",
     "Public Function Probe() As String\n"
     "    Probe = modSimReport.SimReportDerivedStatus()\n"
     "End Function", "ownership"),
    # A FINGERPRINT PROCEDURE NAMED OUTSIDE ITS TWO OWNERS.
    ("a third module names a fingerprint procedure", "modSimEngine",
     "Public Function Probe() As String\n"
     "    Probe = SimFpResultDigest()\n"
     "End Function", "ownership"),
    # THE MACHINE SHEET NAMED BY THE PRESENTATION ADAPTER.
    ("the adapter names the machine sheet", "modResultsState",
     "Public Function Probe() As String\n"
     '    Probe = "_SimData"\n'
     "End Function", "fingerprint"),
])
def test_95_the_procedure_level_detector_refuses_each_ownership_breach(
        name: str, module: str, injected: str, rule: str) -> None:
    """EACH BREACH INJECTED INTO A COPY OF THE MODULE SET IN MEMORY, and run
    against the same two rules the real tree satisfies. Nothing on disk
    changes."""
    import test_phase6_request_fingerprint as controls

    modules = _modules_with(module, injected)
    refused = []
    for rule_name, checker in (("ownership", controls._assert_sim_report_ownership),
                               ("fingerprint",
                                controls._assert_results_state_owns_no_fingerprint)):
        if rule not in ("both", rule_name):
            continue
        try:
            checker(modules)
        except AssertionError as failure:
            refused.append(f"{rule_name}: {failure}")
    assert refused, f"'{name}' survived the detector"


def test_96_the_procedure_level_detector_passes_on_the_real_tree() -> None:
    """SO THE SEVEN REFUSALS ABOVE ARE REFUSALS OF THE INJECTION, not of the
    fixture - and the declared call sets are exactly what the tree does."""
    import re as _re

    import test_phase6_request_fingerprint as controls
    from pccm_builder.vba_source import load_modules

    modules = load_modules([SRC])
    controls._assert_sim_report_ownership(modules)
    controls._assert_results_state_owns_no_fingerprint(modules)
    actual = {}
    for module in modules:
        if module.name == "modSimReport":
            continue
        called = set(_re.findall(r"modSimReport\.(\w+)", module.code))
        if called:
            actual[module.name] = called
    assert actual == controls.SIM_REPORT_CALLERS, {
        k: sorted(v) for k, v in actual.items()}
    assert controls.SIM_REPORT_CALLERS["modResultsState"] == {PURE_OWNER}, (
        "the worksheet adapter's permitted dependency is no longer the pure evaluator")
